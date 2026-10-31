"""Build the Agentic AI course presentation (.pptx, 16:9).

Same visual system as build_ppt.py (stream-processing deck).
Writes: presentation/Wildfire_Agentic_AI.pptx
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation" / "Wildfire_Agentic_AI.pptx"

INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x89, 0x87, 0x81)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
RED = RGBColor(0xE3, 0x49, 0x48)
AQUA = RGBColor(0x1B, 0xAF, 0x7A)
VIOLET = RGBColor(0x4A, 0x3A, 0xA7)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
PANEL = RGBColor(0xF6, 0xF5, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x11, 0x16, 0x22)
FONT = "Calibri"

SW, SH = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]
slide_no = [0]


def slide():
    s = prs.slides.add_slide(BLANK)
    slide_no[0] += 1
    return s


def box(s, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill if fill else WHITE
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.0, space_after=4):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [(runs, {})]
    for i, (t, ov) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ov.get("align", align)
        p.space_after = Pt(ov.get("space_after", space_after))
        p.line_spacing = ov.get("spacing", spacing)
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = ov.get("font", FONT)
        f.size = Pt(ov.get("size", size))
        f.bold = ov.get("bold", bold)
        f.italic = ov.get("italic", False)
        f.color.rgb = ov.get("color", color)
    return tb


def header(s, kicker, title, accent=VIOLET):
    box(s, 0, 0, SW, Inches(0.10), fill=accent)
    text(s, Inches(0.55), Inches(0.32), Inches(12.2), Inches(0.3),
         kicker.upper(), size=12, color=accent, bold=True)
    text(s, Inches(0.55), Inches(0.62), Inches(12.2), Inches(0.75),
         title, size=27, bold=True, color=INK)
    text(s, Inches(12.35), Inches(7.06), Inches(0.6), Inches(0.3),
         f"{slide_no[0]:02d}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    text(s, Inches(0.55), Inches(7.06), Inches(8), Inches(0.3),
         "Agentic Wildfire Intelligence | Agentic AI", size=9, color=MUTED)


def bullets(s, x, y, w, h, items, size=15, gap=8, color=INK):
    runs = []
    for it in items:
        if isinstance(it, tuple):
            head, rest = it
            runs.append((f"▪  {head}", {"bold": True, "size": size, "color": color, "space_after": 2}))
            runs.append((f"     {rest}", {"size": size - 1, "color": INK2, "space_after": gap}))
        else:
            runs.append((f"▪  {it}", {"size": size, "color": color, "space_after": gap}))
    text(s, x, y, w, h, runs)


def table(s, x, y, w, h, rows, col_w=None, header_fill=VIOLET,
          fs=11.5, header_fs=12, row_h=Inches(0.34)):
    n_r, n_c = len(rows), len(rows[0])
    gfx = s.shapes.add_table(n_r, n_c, x, y, w, h)
    tbl = gfx.table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = row_h if ri else Inches(0.38)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if ri == 0 else (WHITE if ri % 2 else PANEL)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.name = FONT
            r.font.size = Pt(header_fs if ri == 0 else fs)
            r.font.bold = ri == 0
            r.font.color.rgb = WHITE if ri == 0 else INK
    return tbl


def arrow(s, x1, y1, x2, y2, color=GREY, weight=1.75):
    from pptx.enum.shapes import MSO_CONNECTOR
    conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = Pt(weight)
    return conn


# ============================================================ 01 TITLE
s = slide()
box(s, 0, 0, SW, SH, fill=DARK)
box(s, 0, Inches(4.55), SW, Pt(3), fill=ORANGE)
text(s, Inches(0.9), Inches(1.55), Inches(11.5), Inches(0.4),
     "AGENTIC AI  •  COURSE PROJECT PRESENTATION", size=14, color=ORANGE, bold=True)
text(s, Inches(0.9), Inches(2.0), Inches(11.6), Inches(1.7),
     "Agentic Wildfire Intelligence", size=44, bold=True, color=WHITE)
text(s, Inches(0.9), Inches(3.35), Inches(11.6), Inches(1.0),
     "A multi-agent system that plans, verifies and acts on top of a real-time\nwildfire detection pipeline — coordinator, worker agents, critic and guardrails",
     size=17, color=RGBColor(0xC3, 0xC2, 0xB7))
text(s, Inches(0.9), Inches(4.95), Inches(11), Inches(1.4), [
    ("Presented by:  Pratham Mallya", {"size": 15, "color": WHITE, "space_after": 6}),
    ("Guide:  <Guide Name>   •   Department of <CSE/AI-ML>", {"size": 13, "color": RGBColor(0x8B, 0x93, 0xA6)}),
    ("October 2026", {"size": 13, "color": RGBColor(0x8B, 0x93, 0xA6)}),
])

# ============================================================ 02 AGENDA
s = slide()
header(s, "Agenda", "Presentation Outline")
items = [
    ("01", "Introduction & Motivation", "Why a detection pipeline alone is not enough", BLUE),
    ("02", "Agentic AI Concepts Applied", "Course concepts → concrete implementation", VIOLET),
    ("03", "System Context", "The real-time wildfire pipeline the agents sit on", AQUA),
    ("04", "Multi-Agent Architecture", "Coordinator / worker / critic hub-and-spoke design", ORANGE),
    ("05", "The Agents", "Planner-delegator, four specialist workers, user-facing chat", ORANGE),
    ("06", "Guardrails, Orchestration & Live Demo", "Tool allowlists, approval gates, critic verification, real runs", RED),
    ("07", "Evaluation & Work Completed", "Agent-level metrics, limitations, roadmap", GREY),
]
y = Inches(1.75)
for num, t, sub, c in items:
    box(s, Inches(0.9), y, Inches(0.72), Inches(0.62), fill=c)
    text(s, Inches(0.9), y + Inches(0.10), Inches(0.72), Inches(0.4),
         num, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, Inches(1.85), y + Inches(0.02), Inches(10.6), Inches(0.35),
         t, size=17, bold=True, color=INK)
    text(s, Inches(1.85), y + Inches(0.36), Inches(10.6), Inches(0.3),
         sub, size=12, color=INK2)
    y += Inches(0.74)

# ============================================================ 03 INTRO 1
s = slide()
header(s, "01 Introduction", "From a Model That Predicts to a System That Reasons and Acts", accent=BLUE)
bullets(s, Inches(0.55), Inches(1.7), Inches(7.2), Inches(4.9), [
    ("The pipeline detects — but does not understand",
     "Our streaming system scores every grid-cell reading and emits tiered alerts (WATCH / WARNING / EMERGENCY). It cannot explain, prioritize or act on them."),
    ("Alert volume overwhelms humans",
     "44,000+ WATCH alerts accumulated in Elasticsearch — no operator can manually verify which ones matter, where, or why now."),
    ("Real questions are multi-step",
     "\"Has this location burned before, and does today's alert matter?\" requires planning: query alerts → geo-lookup history → weigh severity → draft a response."),
    ("Actions need judgment and permission",
     "Paging a district officer is irreversible and costly — an autonomous system must draft, verify and ask before it acts."),
], size=14.5)
box(s, Inches(8.15), Inches(1.75), Inches(4.6), Inches(4.85), fill=PANEL)
text(s, Inches(8.45), Inches(1.95), Inches(4.0), Inches(0.4),
     "THE GAP AGENTS FILL", size=12, bold=True, color=VIOLET)
text(s, Inches(8.45), Inches(2.4), Inches(4.0), Inches(4.0), [
    ("Before (pipeline only)", {"bold": True, "size": 13.5, "space_after": 2}),
    ("Events → ML score → alert → dashboard. A human must watch, interpret, investigate and decide — continuously.", {"size": 12.5, "color": INK2, "space_after": 12}),
    ("After (agentic layer)", {"bold": True, "size": 13.5, "space_after": 2}),
    ("LLM agents that plan over the pipeline's tools, verify claims against Elasticsearch evidence, check fire history, and draft dispatches — with a human gate on irreversible actions.", {"size": 12.5, "color": INK2, "space_after": 12}),
    ("Key idea", {"bold": True, "size": 13.5, "space_after": 2}),
    ("The pipeline stays the source of truth; agents reason over it, never around it.", {"size": 12.5, "color": INK2}),
])

# ============================================================ 04 INTRO 2
s = slide()
header(s, "01 Introduction", "Passive Pipeline vs. Agentic Layer", accent=BLUE)
table(s, Inches(0.55), Inches(1.8), Inches(12.2), Inches(3.1), [
    ["Question", "Without agents", "With the agentic layer"],
    ["Which of 44k alerts matter?", "Operator scrolls a dashboard", "Alert Verifier cross-checks ES evidence and ranks by tier/probability"],
    ["Did this area burn before?", "Manual GIS / archive lookup", "Fire History agent runs a geo-query and scores recurrence risk"],
    ["Is the pipeline itself healthy?", "Read Prometheus graphs", "Situation agent summarizes latency, throughput, tier counts in words"],
    ["Notify responders", "Hand-written message", "Dispatcher drafts from verified evidence; sends only after approval"],
    ["Trust the answer?", "—", "Critic checks every claim against tool evidence before answering"],
], col_w=[3.0, 3.7, 5.5], fs=11.5, row_h=Inches(0.5))
text(s, Inches(0.55), Inches(5.35), Inches(12.2), Inches(1.5), [
    ("Agents do not replace the pipeline — they make it operable.", {"bold": True, "size": 14.5, "space_after": 4}),
    ("Every agent capability is grounded in an existing system endpoint (Elasticsearch, the dashboard API, the notifier webhook). The LLM supplies planning, language and judgment; the system supplies facts.",
     {"size": 13.5, "color": INK2}),
])

# ============================================================ 05 CONCEPTS
s = slide()
header(s, "02 Agentic AI Concepts", "Course Concepts → Where Each One Lives in the System")
concepts = [
    ("Tool use / function calling", "7 allowlisted tools over ES, the dashboard API and the notifier webhook; the model requests a tool by name, the Toolbelt executes it — never eval of model output.", BLUE),
    ("Planning & decomposition", "Coordinator's PLAN phase turns a user goal into typed JSON subtasks for the worker registry, with a needs_approval flag.", VIOLET),
    ("Multi-agent orchestration", "Hub-and-spoke: 4 specialist workers delegated by the coordinator, run in parallel threads; workers never call each other — every run is a traceable tree.", ORANGE),
    ("ReAct loop (bounded)", "Each worker interleaves reasoning and tool calls, capped at 4 rounds, tool results truncated to fit small-model contexts.", AQUA),
    ("Reflection / self-critique", "A Critic pass re-checks the draft answer against raw tool evidence and flags unsupported claims; the answer is regenerated if ungrounded.", RED),
    ("Human-in-the-loop", "The only side-effecting tool (send_dispatch) refuses to run without approved=true — drafts are free, sends are gated.", GREY),
]
y = Inches(1.68)
for name, desc, c in concepts:
    box(s, Inches(0.55), y, Inches(0.09), Inches(0.78), fill=c)
    text(s, Inches(0.85), y, Inches(3.4), Inches(0.75), name, size=14, bold=True, color=INK)
    text(s, Inches(4.4), y + Inches(0.02), Inches(8.35), Inches(0.8), desc, size=12, color=INK2)
    y += Inches(0.9)

# ============================================================ 06 SYSTEM CONTEXT
s = slide()
header(s, "03 System Context", "What the Agents Sit On: The Streaming Backbone (Recap)", accent=AQUA)
stages = [
    ("Kafka topics", "raw → features → predictions → alerts", BLUE),
    ("Stream processing", "validate · dedup · windows · features", AQUA),
    ("ML detector", "RandomForest fire probability + tiers", ORANGE),
    ("Elasticsearch", "wildfire-alerts index (44k+ docs)", VIOLET),
    ("Dashboard + notifier", ":9991 API · SSE globe · webhook", RED),
]
x = Inches(0.55)
for i, (name, sub, c) in enumerate(stages):
    box(s, x, Inches(1.9), Inches(2.28), Inches(1.35), fill=PANEL, line=c, line_w=1.5)
    text(s, x + Inches(0.12), Inches(2.05), Inches(2.05), Inches(0.4), name, size=13, bold=True, color=c)
    text(s, x + Inches(0.12), Inches(2.5), Inches(2.05), Inches(0.7), sub, size=10.5, color=INK2)
    if i < 4:
        arrow(s, x + Inches(2.28), Inches(2.57), x + Inches(2.52), Inches(2.57), color=GREY)
    x += Inches(2.52)
box(s, Inches(0.55), Inches(3.75), Inches(12.2), Inches(0.02), fill=RGBColor(0xE1, 0xE0, 0xD9))
text(s, Inches(0.55), Inches(3.9), Inches(12.2), Inches(0.35),
     "AGENTIC LAYER  (this project)  —  reads and acts through the interfaces below", size=12.5, bold=True, color=VIOLET)
rows = [
    ["System interface", "What it exposes", "Which agent uses it"],
    ["Elasticsearch wildfire-alerts", "Full-text + tier search; lat/lon range geo queries", "Alert Verifier · Fire History"],
    ["Dashboard API (:9991)", "/api/alerts · /api/predictions · /api/stats · /api/search", "Situation · Alert Verifier"],
    ["Notifier webhook", "Emergency dispatch channel", "Dispatcher (gated)"],
    ["Pipeline metrics", "Latency p50/p95/p99, throughput, tier counts", "Situation"],
]
table(s, Inches(0.55), Inches(4.35), Inches(12.2), Inches(2.3), rows,
      col_w=[3.3, 5.4, 3.5], fs=11.5, header_fill=AQUA, row_h=Inches(0.42))

# ============================================================ 07 ARCHITECTURE
s = slide()
header(s, "04 Architecture", "Hub-and-Spoke Multi-Agent Architecture", accent=ORANGE)
# user + chat
box(s, Inches(0.55), Inches(1.75), Inches(2.5), Inches(1.0), fill=PANEL, line=BLUE, line_w=1.5)
text(s, Inches(0.67), Inches(1.86), Inches(2.3), Inches(0.35), "User / Operator", size=12.5, bold=True, color=BLUE)
text(s, Inches(0.67), Inches(2.22), Inches(2.3), Inches(0.4), "natural-language goal", size=10, color=INK2)
box(s, Inches(0.55), Inches(3.15), Inches(2.5), Inches(1.0), fill=BLUE)
text(s, Inches(0.67), Inches(3.26), Inches(2.3), Inches(0.35), "Chat Agent", size=12.5, bold=True, color=WHITE)
text(s, Inches(0.67), Inches(3.62), Inches(2.3), Inches(0.4), "user-facing Q&A (CLI/REPL)", size=10, color=WHITE)
# coordinator
box(s, Inches(4.35), Inches(2.35), Inches(3.4), Inches(1.9), fill=VIOLET)
text(s, Inches(4.5), Inches(2.46), Inches(3.1), Inches(0.35), "Coordinator Agent", size=14, bold=True, color=WHITE)
text(s, Inches(4.5), Inches(2.84), Inches(3.15), Inches(1.3),
     "PLAN → DELEGATE → VERIFY → SYNTHESIZE\ngpt-oss-120b · JSON plans · approval flag\nparallel delegation (thread pool)", size=10.5, color=WHITE)
# workers
workers = [
    ("Alert Verifier", "search_alerts · recent_alerts\nlatest_predictions", BLUE),
    ("Fire History", "fire_history (geo) · search_alerts", AQUA),
    ("Situation", "pipeline_stats · latest_predictions\nrecent_alerts", ORANGE),
    ("Dispatcher", "draft_dispatch · send_dispatch (gated)", RED),
]
y = Inches(1.55)
for name, tools_, c in workers:
    box(s, Inches(9.05), y, Inches(3.75), Inches(1.06), fill=PANEL, line=c, line_w=1.5)
    text(s, Inches(9.2), y + Inches(0.08), Inches(3.5), Inches(0.35), name, size=12.5, bold=True, color=c)
    text(s, Inches(9.2), y + Inches(0.44), Inches(3.5), Inches(0.55), tools_, size=9.5, color=INK2)
    y += Inches(1.28)
# critic
box(s, Inches(4.35), Inches(4.75), Inches(3.4), Inches(0.95), fill=RED)
text(s, Inches(4.5), Inches(4.85), Inches(3.1), Inches(0.35), "Critic (reflection)", size=12.5, bold=True, color=WHITE)
text(s, Inches(4.5), Inches(5.2), Inches(3.1), Inches(0.4), "every claim must trace to tool evidence", size=10, color=WHITE)
# arrows
arrow(s, Inches(1.8), Inches(2.75), Inches(1.8), Inches(3.15))
arrow(s, Inches(3.05), Inches(3.65), Inches(4.35), Inches(3.35))
arrow(s, Inches(7.75), Inches(2.95), Inches(9.05), Inches(2.1))
arrow(s, Inches(7.75), Inches(3.2), Inches(9.05), Inches(3.35))
arrow(s, Inches(7.75), Inches(3.5), Inches(9.05), Inches(4.6))
arrow(s, Inches(7.75), Inches(3.75), Inches(9.05), Inches(5.85))
arrow(s, Inches(6.05), Inches(4.25), Inches(6.05), Inches(4.75), color=RED)
arrow(s, Inches(6.05), Inches(5.7), Inches(6.05), Inches(6.15), color=RED)
text(s, Inches(4.35), Inches(6.2), Inches(4.6), Inches(0.4),
     [("ungrounded → regenerate", {"size": 10, "italic": True, "color": RED})])
# infra strip
box(s, Inches(9.05), Inches(6.45), Inches(3.75), Inches(0.55), fill=DARK)
text(s, Inches(9.15), Inches(6.54), Inches(3.6), Inches(0.4),
     "Tools → ES :9200 · API :9991 · webhook", size=10, color=WHITE)
text(s, Inches(0.55), Inches(6.45), Inches(7.6), Inches(0.55),
     "Workers: gpt-oss-20b (separate rate-limit bucket)  •  every prompt & tool I/O logged to logs/agent_traces.jsonl",
     size=10.5, color=INK2)

# ============================================================ 08 ROSTER
s = slide()
header(s, "05 The Agents", "Agent Roster: Roles, Tools and Models", accent=ORANGE)
table(s, Inches(0.55), Inches(1.75), Inches(12.2), Inches(3.4), [
    ["Agent", "Role (prompt-defined)", "Allowlisted tools", "Model"],
    ["ChatAgent", "User-facing Q&A; keeps conversational context, hands each turn to the coordinator", "— (via coordinator)", "gpt-oss-120b"],
    ["CoordinatorAgent", "Planner + delegator: JSON plan, parallel delegation, approval flag, final synthesis", "all (via workers)", "gpt-oss-120b"],
    ["AlertVerifierAgent", "Confirms/refutes fire claims against ES alert + prediction evidence", "search_alerts, recent_alerts, latest_predictions", "gpt-oss-20b"],
    ["FireHistoryAgent", "\"Has this location burned before?\" — geo query + recurrence scoring", "fire_history, search_alerts", "gpt-oss-20b"],
    ["SituationAgent", "Current conditions + pipeline health in plain language", "pipeline_stats, latest_predictions, recent_alerts", "gpt-oss-20b"],
    ["DispatcherAgent", "Drafts emergency notifications; sends only with approval", "draft_dispatch, send_dispatch", "gpt-oss-20b"],
    ["Critic", "Verifies the draft answer is grounded in tool evidence", "— (reads evidence)", "gpt-oss-120b"],
], col_w=[2.4, 4.8, 3.2, 1.8], fs=10.5, row_h=Inches(0.46))
text(s, Inches(0.55), Inches(5.5), Inches(12.2), Inches(1.3), [
    ("Why two models?", {"bold": True, "size": 13.5, "space_after": 3}),
    ("Groq rate limits are per-model tokens/minute. Planning/verification (large contexts) runs on gpt-oss-120b; workers (small, frequent calls) run on gpt-oss-20b — two independent rate-limit buckets, so coordinator traffic can never starve worker traffic. Also a cost/latency split: big model where reasoning matters, small where it doesn't.",
     {"size": 12, "color": INK2}),
])

# ============================================================ 09 COORDINATOR
s = slide()
header(s, "05 The Agents", "Coordinator: Plan → Delegate → Verify → Synthesize", accent=ORANGE)
phases = [
    ("1 · PLAN", "LLM decomposes the goal into typed subtasks for the worker registry; sets needs_approval if a send/notify is requested.", VIOLET),
    ("2 · DELEGATE", "Each subtask runs on its worker in a thread pool (parallel). Unknown workers are rejected; a failed worker degrades to an error note, not a crash.", BLUE),
    ("3 · VERIFY", "Critic receives the draft + raw tool evidence; any number, tier, location or timestamp not in the evidence is flagged.", RED),
    ("4 · SYNTHESIZE", "Final answer: Situation (verified facts) → Assessment → Recommended actions; tool failures disclosed, not hidden.", AQUA),
]
y = Inches(1.7)
for name, desc, c in phases:
    box(s, Inches(0.55), y, Inches(0.09), Inches(1.0), fill=c)
    text(s, Inches(0.85), y, Inches(2.1), Inches(0.4), name, size=14, bold=True, color=c)
    text(s, Inches(3.0), y + Inches(0.02), Inches(9.7), Inches(1.0), desc, size=12.5, color=INK2)
    y += Inches(1.14)
box(s, Inches(0.55), Inches(6.15), Inches(12.2), Inches(0.85), fill=DARK)
text(s, Inches(0.8), Inches(6.27), Inches(11.8), Inches(0.65),
     [("Real plan produced for \"Has 24.07, 80.90 burned before?\":", {"size": 11, "color": RGBColor(0x8B, 0x93, 0xA6), "space_after": 2}),
      ('{"subtasks": [{"worker": "fire_history", "task": "Check earlier fires at lat 24.07, lon 80.90 and summarize implications"}], "needs_approval": false}',
       {"size": 11, "color": WHITE, "font": "Consolas"})])

# ============================================================ 10 VERIFIER
s = slide()
header(s, "05 The Agents", "Worker Spotlight: Alert Verifier (the Elasticsearch Checker)", accent=BLUE)
bullets(s, Inches(0.55), Inches(1.7), Inches(7.2), Inches(4.7), [
    ("Job", "Given any claim about fire activity, verify it against Elasticsearch — confirm, contradict, or declare unverifiable."),
    ("Tools", "search_alerts (full-text + tier), recent_alerts, latest_predictions — read-only."),
    ("Method", "Cross-checks alert documents against model predictions: a tier without matching high probability is flagged as suspicious."),
    ("Honesty rule", "If tools error or the index is down, it must say so — the system prompt forbids inventing alert data."),
], size=14)
box(s, Inches(8.15), Inches(1.75), Inches(4.6), Inches(4.85), fill=PANEL)
text(s, Inches(8.45), Inches(1.95), Inches(4.0), Inches(0.4),
     "LIVE OUTPUT (30 OCT 2026)", size=11.5, bold=True, color=BLUE)
text(s, Inches(8.45), Inches(2.4), Inches(4.0), Inches(3.9), [
    ("Query: \"Any active WARNING or EMERGENCY alerts right now?\"", {"size": 11.5, "italic": True, "color": INK2, "space_after": 8}),
    ("\"No active wildfire alerts with severity WARNING or EMERGENCY were found in Elasticsearch. The search returned zero results for both tiers.\"", {"size": 12, "color": INK, "space_after": 10}),
    ("Verified against 44,473 indexed alert documents — all currently WATCH tier.", {"size": 11, "color": MUTED}),
])

# ============================================================ 11 HISTORY
s = slide()
header(s, "05 The Agents", "Worker Spotlight: Fire History (Has This Place Burned Before?)", accent=AQUA)
bullets(s, Inches(0.55), Inches(1.7), Inches(7.2), Inches(4.7), [
    ("Job", "Geo-lookup of prior alerts around any lat/lon, then a plain-language recurrence verdict for planners."),
    ("Geo query", "The index stores lat/lon floats (not geo_point), so the tool runs a bounding-box range query server-side, then refines to a true radius with haversine client-side."),
    ("Recurrence rule", "Severity-weighted score — WATCH×1, WARNING×3, EMERGENCY×6 — with an automatic 'high' if any prior EMERGENCY exists."),
    ("Why it matters", "A WATCH in a never-burned grid cell is noise; the same WATCH in a cell with 200 prior alerts is a pattern — context the raw pipeline cannot supply."),
], size=14)
box(s, Inches(8.15), Inches(1.75), Inches(4.6), Inches(4.85), fill=PANEL)
text(s, Inches(8.45), Inches(1.95), Inches(4.0), Inches(0.4),
     "LIVE OUTPUT — 24.07°N, 80.90°E, 25 KM", size=11.5, bold=True, color=AQUA)
text(s, Inches(8.45), Inches(2.45), Inches(4.0), Inches(3.9), [
    ("200 prior alerts", {"size": 20, "bold": True, "color": ORANGE, "space_after": 2}),
    ("all WATCH tier · most recent 30 Oct 2026, 0.5 km away", {"size": 11.5, "color": INK2, "space_after": 10}),
    ("Recurrence score 200 → level: HIGH", {"size": 13.5, "bold": True, "color": RED, "space_after": 8}),
    ("\"Frequent ignition-favourable conditions… elevated risk of fire starts when dry fuels, wind and heat align.\"", {"size": 11.5, "italic": True, "color": INK}),
])

# ============================================================ 12 SITUATION + DISPATCHER
s = slide()
header(s, "05 The Agents", "Worker Spotlight: Situation & Dispatcher (the Action Agents)", accent=RED)
bullets(s, Inches(0.55), Inches(1.7), Inches(6.6), Inches(2.3), [
    ("SituationAgent", "Reads /api/stats, latest predictions and recent alerts; separates the *fire* situation from *system* health — \"no WARNINGs active; p95 latency 219 ms; 50 events/s; ES up\"."),
    ("DispatcherAgent", "Turns verified findings into a dispatch: tier, region, key evidence, one recommended action."),
], size=13.5)
box(s, Inches(0.55), Inches(4.15), Inches(6.6), Inches(2.4), fill=PANEL)
text(s, Inches(0.8), Inches(4.3), Inches(6.1), Inches(0.35), "THE APPROVAL GATE (CODE, NOT PROMPT)", size=11.5, bold=True, color=RED)
text(s, Inches(0.8), Inches(4.7), Inches(6.15), Inches(1.75), [
    ("def send_dispatch(message, approved=False):", {"size": 11.5, "font": "Consolas", "color": INK, "space_after": 2}),
    ("    if not approved:", {"size": 11.5, "font": "Consolas", "color": INK, "space_after": 2}),
    ('        return {"sent": False, "refused": True,', {"size": 11.5, "font": "Consolas", "color": INK, "space_after": 2}),
    ('                "reason": "requires approval gate"}', {"size": 11.5, "font": "Consolas", "color": INK}),
])
# approval flow right
steps = [
    ("1", "User asks: \"notify the district officer\"", INK),
    ("2", "Planner sets needs_approval=true", VIOLET),
    ("3", "Dispatcher DRAFTS only — send refused", BLUE),
    ("4", "Coordinator returns draft + ⚠ approval request", ORANGE),
    ("5", "Human says yes → re-run with approved=true", RED),
    ("6", "Webhook fires; audit trail in trace log", AQUA),
]
y = Inches(1.7)
text(s, Inches(7.6), y - Inches(0.1), Inches(5.2), Inches(0.3),
     "HUMAN-IN-THE-LOOP FLOW", size=11.5, bold=True, color=RED)
y += Inches(0.32)
for n, t, c in steps:
    box(s, Inches(7.6), y, Inches(0.42), Inches(0.42), fill=c)
    text(s, Inches(7.6), y + Inches(0.05), Inches(0.42), Inches(0.3),
         n, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, Inches(8.2), y + Inches(0.05), Inches(4.6), Inches(0.5), t, size=11.5, color=INK)
    y += Inches(0.78)

# ============================================================ 13 GUARDRAILS
s = slide()
header(s, "06 Guardrails & Orchestration", "Safety and Reliability Mechanisms", accent=RED)
rows = [
    ["Mechanism", "Implementation", "Failure it prevents"],
    ["Tool allowlist", "Model picks a name; Toolbelt maps it to a function. Anything unlisted → rejected.", "Arbitrary code execution / prompt-injection damage"],
    ["Approval gate", "send_dispatch refuses without approved=true; the flag comes from the human, never the model", "Autonomous paging of real responders"],
    ["Critic verification", "Every quantitative claim must trace to raw tool evidence; ungrounded drafts are regenerated", "Hallucinated alerts presented as fact"],
    ["Bounded loops", "Max 4 tool rounds/worker; tool results truncated to 2,000 chars", "Runaway agents; context overflow (413)"],
    ["Graceful degradation", "Tool errors return {\"unavailable\": true}; a dead worker becomes an error note the answer must disclose", "Silent wrong answers when ES/API is down"],
    ["Full trace logging", "Every prompt, tool call and result appended to logs/agent_traces.jsonl", "Irreproducible behavior; no audit trail"],
    ["Retry + backoff", "429-aware backoff (8/20/40/60 s); per-model rate buckets", "Rate-limit collapse under parallel workers"],
]
table(s, Inches(0.55), Inches(1.75), Inches(12.2), Inches(4.6), rows,
      col_w=[2.3, 5.6, 4.3], fs=11, row_h=Inches(0.56), header_fill=RED)

# ============================================================ 14 ORCHESTRATION SEQUENCE
s = slide()
header(s, "06 Guardrails & Orchestration", "Anatomy of One Run (Live Fire-History Query)", accent=RED)
seq = [
    ("User", "ChatAgent", "\"Has the area around 24.07, 80.90 had earlier fires? What does that imply?\"", BLUE),
    ("ChatAgent", "Coordinator", "goal handed off for planning", BLUE),
    ("Coordinator", "LLM (120b)", "PLAN → 1 subtask: fire_history worker, needs_approval=false", VIOLET),
    ("Coordinator", "FireHistoryAgent", "DELEGATE (thread pool)", ORANGE),
    ("FireHistoryAgent", "Toolbelt", "fire_history(lat=24.07, lon=80.90, radius_km=25)", AQUA),
    ("Toolbelt", "Elasticsearch", "bounding-box range query → 200 hits → haversine refine", AQUA),
    ("FireHistoryAgent", "Coordinator", "findings: 200 WATCH, score 200 → HIGH recurrence", ORANGE),
    ("Coordinator", "Critic (120b)", "VERIFY draft vs evidence → 1 issue caught (latest-alert timestamp)", RED),
    ("Coordinator", "User", "SYNTHESIZE → Situation / Assessment / Recommended actions", INK),
]
y = Inches(1.62)
box(s, Inches(0.55), y - Inches(0.05), Inches(2.2), Inches(0.35), fill=DARK)
text(s, Inches(0.55), y, Inches(2.2), Inches(0.3), "FROM → TO", size=10.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
text(s, Inches(3.0), y, Inches(9.7), Inches(0.3), "MESSAGE", size=10.5, bold=True, color=MUTED)
y += Inches(0.42)
for frm, to, msg, c in seq:
    box(s, Inches(0.55), y, Inches(2.2), Inches(0.5), fill=PANEL, line=c, line_w=1.0)
    text(s, Inches(0.55), y + Inches(0.11), Inches(2.2), Inches(0.3),
         f"{frm} → {to}", size=10, bold=True, color=c, align=PP_ALIGN.CENTER)
    text(s, Inches(3.0), y + Inches(0.08), Inches(9.7), Inches(0.5), msg, size=11, color=INK)
    y += Inches(0.585)

# ============================================================ 15 DEMO
s = slide()
header(s, "06 Guardrails & Orchestration", "Live Demo Results (Real Runs, 30–31 Oct 2026)", accent=RED)
box(s, Inches(0.55), Inches(1.7), Inches(6.0), Inches(4.9), fill=PANEL)
text(s, Inches(0.8), Inches(1.85), Inches(5.5), Inches(0.35),
     "SCENARIO A — LOCATION HISTORY", size=11.5, bold=True, color=AQUA)
text(s, Inches(0.8), Inches(2.25), Inches(5.55), Inches(4.2), [
    ("Q: \"Has 24.07, 80.90 had earlier fires? What does that imply?\"", {"size": 11, "italic": True, "color": INK2, "space_after": 6}),
    ("Plan: 1 subtask → fire_history worker", {"size": 11, "color": INK, "space_after": 4}),
    ("Evidence: 200 prior alerts in 25 km, all WATCH, latest 0.5 km away; recurrence HIGH", {"size": 11, "color": INK, "space_after": 4}),
    ("Critic: caught a wrong \"most recent alert\" timestamp in the draft → final answer corrected", {"size": 11, "color": RED, "space_after": 4}),
    ("Answer: Situation → Assessment (\"elevated risk of fire starts\") → 5 recommended actions", {"size": 11, "color": INK, "space_after": 4}),
    ("End-to-end: ~35 s, 4 LLM calls, 1 ES query", {"size": 10.5, "color": MUTED}),
])
box(s, Inches(6.85), Inches(1.7), Inches(5.9), Inches(4.9), fill=PANEL)
text(s, Inches(7.1), Inches(1.85), Inches(5.4), Inches(0.35),
     "SCENARIO B — GRACEFUL DEGRADATION", size=11.5, bold=True, color=RED)
text(s, Inches(7.1), Inches(2.25), Inches(5.45), Inches(4.2), [
    ("Q: \"Any WARNING/EMERGENCY alerts? Is the pipeline healthy?\"", {"size": 11, "italic": True, "color": INK2, "space_after": 6}),
    ("Plan: 2 subtasks → alert_verifier + situation (parallel)", {"size": 11, "color": INK, "space_after": 4}),
    ("Verifier: \"Zero WARNING/EMERGENCY in ES\" ✓", {"size": 11, "color": INK, "space_after": 4}),
    ("Situation worker: Groq API context overflow → worker failed after backoff", {"size": 11, "color": RED, "space_after": 4}),
    ("Final answer disclosed the failure instead of hiding it: \"The situation worker did not return data… this does not affect the alert status.\"", {"size": 11, "color": INK, "space_after": 4}),
    ("→ the guardrail worked: partial truth, honestly labelled", {"size": 11, "bold": True, "color": AQUA}),
])

# ============================================================ 16 IMPLEMENTATION
s = slide()
header(s, "06 Guardrails & Orchestration", "Implementation Stack", accent=GREY)
rows = [
    ["Component", "Choice", "Why"],
    ["LLM backend", "Groq API (OpenAI-compatible chat completions)", "Fast inference, generous free tier, native tool calling + JSON mode"],
    ["Coordinator / Critic", "openai/gpt-oss-120b", "Largest context for plans, synthesis, evidence checking"],
    ["Workers", "openai/gpt-oss-20b", "Separate rate-limit bucket; cheap, fast, tool-capable"],
    ["Agent framework", "None — ~450 lines of plain Python", "Full control of the loop; every design decision visible for the course"],
    ["Tools", "requests → ES :9200, dashboard API :9991, notifier webhook", "Zero new infrastructure — agents ride the existing system"],
    ["Concurrency", "ThreadPoolExecutor over subtasks", "Workers run in parallel; wall-clock ≈ slowest worker"],
    ["Observability", "logs/agent_traces.jsonl (every prompt + tool I/O)", "Reproducibility, audit, debugging"],
    ["Resilience", "Timeout 30 s · 4 retries · 429-aware backoff · degradation", "Network calls are never trusted"],
]
table(s, Inches(0.55), Inches(1.75), Inches(12.2), Inches(3.6), rows,
      col_w=[2.5, 4.4, 5.3], fs=11, row_h=Inches(0.44), header_fill=GREY)
box(s, Inches(0.55), Inches(5.6), Inches(12.2), Inches(1.15), fill=DARK)
text(s, Inches(0.8), Inches(5.72), Inches(11.7), Inches(0.9), [
    ("Run it:", {"size": 12, "bold": True, "color": RGBColor(0x8B, 0x93, 0xA6), "space_after": 3}),
    ("export GROQ_API_KEY=…    python -m agents.chat \"has 24.07, 80.90 burned before?\"    python -m agents.demo",
     {"size": 12.5, "font": "Consolas", "color": WHITE}),
])

# ============================================================ 17 EVALUATION
s = slide()
header(s, "07 Evaluation & Status", "How We Evaluate the Agent Layer", accent=GREY)
bullets(s, Inches(0.55), Inches(1.7), Inches(7.2), Inches(4.8), [
    ("Task completion", "Did the plan cover the goal? Measured over a scenario set (history lookup, situation report, verification, dispatch draft)."),
    ("Groundedness", "Critic pass rate: fraction of answers whose claims all trace to tool evidence — target 100%, critic catches the rest."),
    ("Tool-call correctness", "Valid tool names, valid argument types, no allowlist violations (0 observed)."),
    ("Latency & cost", "~30–45 s end-to-end per complex query; 3–8 LLM calls; two-model split keeps the 8k TPM free tier viable."),
    ("Honest limitations", "Small worker model can overflow context on large tool payloads (mitigated by truncation); plans occasionally under-decompose compound questions; ES geo is bounding-box + haversine, not a true geo_point index."),
], size=13.5)
box(s, Inches(8.15), Inches(1.75), Inches(4.6), Inches(4.85), fill=PANEL)
text(s, Inches(8.45), Inches(1.95), Inches(4.0), Inches(0.4),
     "MEASURED SO FAR", size=12, bold=True, color=GREY)
text(s, Inches(8.45), Inches(2.4), Inches(4.0), Inches(4.0), [
    ("3/3", {"size": 19, "bold": True, "color": AQUA, "space_after": 1}),
    ("canonical scenarios completed end-to-end against the live pipeline", {"size": 11.5, "color": INK2, "space_after": 10}),
    ("1", {"size": 19, "bold": True, "color": RED, "space_after": 1}),
    ("hallucinated timestamp caught by the critic before reaching the user", {"size": 11.5, "color": INK2, "space_after": 10}),
    ("0", {"size": 19, "bold": True, "color": VIOLET, "space_after": 1}),
    ("ungated sends — approval gate refused every unauthorized dispatch attempt", {"size": 11.5, "color": INK2}),
])

# ============================================================ 18 STATUS + FUTURE
s = slide()
header(s, "07 Evaluation & Status", "Work Completed & Roadmap", accent=GREY)
done = [
    "Multi-agent architecture: coordinator (plan/delegate/verify/synthesize) + 4 workers + critic + chat CLI",
    "7 allowlisted tools over Elasticsearch, dashboard API and notifier webhook",
    "Approval gate on the only side-effecting action (dispatch)",
    "Full trace logging; 429-aware retry; bounded ReAct loops; graceful degradation",
    "Live validation against the running pipeline (44k+ indexed alerts)",
]
future = [
    "Event-driven triggering: agent auto-investigates each EMERGENCY alert from the Kafka topic (no human prompt needed)",
    "Memory: per-location case files so recurring areas accumulate context across runs",
    "Richer workers: weather-forecast tool, evacuation-route planner, resource estimator",
    "Formal eval harness: scenario suite with groundedness scoring in CI",
    "Web UI: embed the chat agent into the 3D dashboard",
]
text(s, Inches(0.55), Inches(1.65), Inches(6.0), Inches(0.35),
     "COMPLETED", size=13, bold=True, color=AQUA)
bullets(s, Inches(0.55), Inches(2.05), Inches(6.0), Inches(4.5),
        [f"{d}" for d in done], size=12, gap=10)
text(s, Inches(7.0), Inches(1.65), Inches(5.8), Inches(0.35),
     "ROADMAP", size=13, bold=True, color=ORANGE)
bullets(s, Inches(7.0), Inches(2.05), Inches(5.8), Inches(4.5),
        [f"{f}" for f in future], size=12, gap=10)

# ============================================================ 19 THANKS
s = slide()
box(s, 0, 0, SW, SH, fill=DARK)
box(s, 0, Inches(4.55), SW, Pt(3), fill=ORANGE)
text(s, Inches(0.9), Inches(2.5), Inches(11.6), Inches(1.0),
     "Thank You", size=44, bold=True, color=WHITE)
text(s, Inches(0.9), Inches(3.6), Inches(11.6), Inches(0.6),
     "Questions & live demo:  python -m agents.chat", size=17, color=RGBColor(0xC3, 0xC2, 0xB7))
text(s, Inches(0.9), Inches(4.95), Inches(11), Inches(1.0), [
    ("Agentic Wildfire Intelligence — agents that plan, verify and act on a real-time fire pipeline", {"size": 13, "color": RGBColor(0x8B, 0x93, 0xA6), "space_after": 4}),
    ("Pratham Mallya  •  Agentic AI  •  October 2026", {"size": 13, "color": RGBColor(0x8B, 0x93, 0xA6)}),
])

prs.save(str(OUT))
print(f"wrote {OUT} ({slide_no[0]} slides)")
